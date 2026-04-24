"""
Extracción de texto desde diferentes formatos de archivo.
Soporta: PDF, DOCX, PPTX, TXT, MD
"""

import io
import fitz  # PyMuPDF para PDFs
import docx  # python-docx
import pptx  # python-pptx


def extraer_texto(file_content: bytes, filename: str) -> str:
    """
    Extrae texto de un archivo en varios formatos.
    
    Args:
        file_content: Contenido del archivo en bytes
        filename: Nombre del archivo (para detectar formato)
    
    Returns:
        Texto extraído
    
    Raises:
        ValueError: Si el formato no es soportado o hay error en extracción
    """
    ext = filename.split('.')[-1].lower()
    text = ""
    
    try:
        if ext == "pdf":
            text = _extraer_texto_pdf(file_content)
        elif ext == "docx":
            text = _extraer_texto_docx(file_content)
        elif ext == "pptx":
            text = _extraer_texto_pptx(file_content)
        elif ext in ["txt", "md"]:
            text = file_content.decode('utf-8', errors='ignore')
        else:
            raise ValueError(f"Formato no soportado: {ext}")      
    except Exception as e:        
        raise ValueError(f"Error extrayendo texto de '{filename}': {str(e)}")

    return text.strip()


def _extraer_texto_pdf(file_content: bytes) -> str:
    """Extrae texto de archivos PDF usando PyMuPDF"""
    doc = fitz.open(stream=file_content, filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text() + "\n"
    doc.close()
    return text


def _extraer_texto_docx(file_content: bytes) -> str:
    """Extrae texto de archivos DOCX"""
    doc = docx.Document(io.BytesIO(file_content))
    text = "\n".join([p.text for p in doc.paragraphs])
    return text


def _extraer_texto_pptx(file_content: bytes) -> str:
    """Extrae texto de archivos PPTX"""
    prs = pptx.Presentation(io.BytesIO(file_content))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text
